import json
import docx

with open(r"D:\ITA\scratch_pages_out.json", encoding="utf-8") as f:
    data = json.load(f)
sections = data["sections"]
figures = data["figures"]

PATH = r"D:\ITA\report\CropSentinel_Mini_Project_Report.docx"
doc = docx.Document(PATH)

toc = doc.tables[1]
lof = doc.tables[2]

FRONT = {"ABSTRACT": "iii", "LIST OF FIGURES": "vii", "LIST OF ABBREVIATIONS": "viii"}

def set_cell_text(cell, text):
    for p in cell.paragraphs:
        if p.runs:
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(text)

matched, unmatched = 0, 0
for row in toc.rows:
    chap = row.cells[0].text.strip()
    title = row.cells[1].text.strip()
    if title in FRONT:
        set_cell_text(row.cells[2], FRONT[title])
        matched += 1
        continue
    key = f"{chap}|{title}"
    if key in sections and sections[key]:
        set_cell_text(row.cells[2], sections[key])
        matched += 1
    elif title == "REFERENCES":
        key2 = "|REFERENCES"
        if sections.get(key2):
            set_cell_text(row.cells[2], sections[key2])
            matched += 1
        else:
            unmatched += 1
    elif title:
        unmatched += 1

print(f"TOC: matched {matched}, unmatched {unmatched}")

lof_matched = 0
for row in lof.rows[1:]:
    fig_no = row.cells[0].text.replace("Figure", "").strip()
    if fig_no in figures and figures[fig_no]:
        set_cell_text(row.cells[2], figures[fig_no])
        lof_matched += 1
print(f"LOF: matched {lof_matched} of {len(lof.rows) - 1}")

doc.save(PATH)
print("patched and saved:", PATH)
